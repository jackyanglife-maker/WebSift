"""DOCX export for parsed records."""

from pathlib import Path
from typing import Iterable
import re

from docx import Document

from database.models import FileLink, Record


def _display_text(value: str) -> str:
    """Render empty text fields with a stable placeholder."""
    return value.strip() or "(empty)"


def _write_markdown_lines(document: Document, content: str) -> None:
    """Render simplified markdown lines into DOCX blocks."""
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s+", line):
            document.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
        else:
            document.add_paragraph(line)


def export_record_to_docx(record: Record, file_links: Iterable[FileLink], output_dir: str = "tmp") -> Path:
    """Write parsed record as .docx file and return file path."""
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    file_path = directory / f"record_{record.id}.docx"

    document = Document()
    document.add_heading(record.title or "Untitled Result", level=1)
    document.add_paragraph(f"URL: {record.url}")
    document.add_paragraph(f"Status: {record.status}")
    document.add_paragraph(f"Page Status: {record.page_status}")

    document.add_heading("Summary", level=2)
    document.add_paragraph(_display_text(record.summary))

    document.add_heading("Content", level=2)
    content = record.content or ""
    if content.strip():
        _write_markdown_lines(document, content)
    else:
        document.add_paragraph("(empty)")

    document.add_heading("Downloadable Files", level=2)
    links = list(file_links)
    if not links:
        document.add_paragraph("None")
    else:
        for item in links:
            document.add_paragraph(
                f"{item.name or item.url} ({item.file_type}): {item.url}",
                style="List Bullet",
            )

    document.save(file_path)
    return file_path
